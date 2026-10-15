from pathlib import Path
from openpyxl import load_workbook
import json

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "source"
OUTPUT = ROOT / "generated"
OUTPUT.mkdir(parents=True, exist_ok=True)

def detect_type(value):
    if value is None:
        return "unknown"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    return "string"

result = {
    "version": "1.0",
    "description": "学校实习管理系统原始Excel字段字典",
    "files": []
}

for file in sorted(SOURCE.glob("*.xlsx")):
    if file.name.startswith("~$"):
        continue

    wb = load_workbook(file, read_only=True, data_only=True)

    file_info = {
        "filename": file.name,
        "sheets": []
    }

    for ws in wb.worksheets:

        # 根据原始模板结构确定表头所在行
        if "暑期高校实习" in file.name:
            header_row = 4
        elif "学科专业数据资源建设指南" in file.name:
            # 该文件包含多张示例表，暂不按普通数据表解析
            file_info["sheets"].append({
                "sheet": ws.title,
                "type": "reference_document",
                "note": "包含7-11、7-12等示例结构，单独作为参考文件"
            })
            continue
        elif ws.title == "实习地区及代码":
            header_row = 0
        else:
            header_row = 1

        fields = []

        if header_row > 0:
            header = next(
                ws.iter_rows(
                    min_row=header_row,
                    max_row=header_row,
                    values_only=True
                ),
                ()
            )

            for index, value in enumerate(header, start=1):
                if value is None:
                    continue

                name = str(value).strip().replace("\n", " ")

                if not name:
                    continue

                fields.append({
                    "column": index,
                    "source_name": name,
                    "data_type": "to_be_confirmed"
                })

        file_info["sheets"].append({
            "sheet": ws.title,
            "header_row": header_row,
            "max_rows": ws.max_row,
            "max_columns": ws.max_column,
            "fields": fields
        })

    wb.close()
    result["files"].append(file_info)

target = OUTPUT / "source_dictionary.json"

target.write_text(
    json.dumps(result, ensure_ascii=False, indent=2),
    encoding="utf-8"
)

print("数据字典生成成功")
print("文件数量：", len(result["files"]))
print("保存位置：", target)

for f in result["files"]:
    print("\n文件：", f["filename"])
    for s in f["sheets"]:
        print(
            "  工作表：",
            s["sheet"],
            "字段数量：",
            len(s.get("fields", []))
        )
