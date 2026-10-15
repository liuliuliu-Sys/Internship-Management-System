from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "source"
OUTPUT = ROOT / "generated"
OUTPUT.mkdir(parents=True, exist_ok=True)

files = sorted(
    f for f in SOURCE.glob("*.xlsx")
    if not f.name.startswith("~$")
)

if len(files) != 6:
    raise SystemExit(f"预期6份Excel，实际找到{len(files)}份，请检查source目录。")

lines = []
total_sheets = 0

for file in files:
    print(f"正在扫描：{file.name}")
    lines.append("=" * 80)
    lines.append(f"文件：{file.name}")

    wb = load_workbook(file, read_only=True, data_only=True)

    for ws in wb.worksheets:
        total_sheets += 1

        lines.append("")
        lines.append(f"工作表：{ws.title}")
        lines.append(f"最大行数：{ws.max_row}")
        lines.append(f"最大列数：{ws.max_column}")
        lines.append("前12行字段结构预览：")

        for number, row in enumerate(
            ws.iter_rows(
                min_row=1,
                max_row=min(ws.max_row or 0, 12),
                max_col=min(ws.max_column or 0, 30),
                values_only=True
            ),
            start=1
        ):
            values = [
                str(value).replace("\n", " ").strip()[:100]
                if value is not None else ""
                for value in row
            ]

            while values and not values[-1]:
                values.pop()

            lines.append(f"第{number}行：" + " | ".join(values))

    wb.close()

report = OUTPUT / "source_scan.txt"
report.write_text("\n".join(lines), encoding="utf-8-sig")

print("")
print("扫描完成！")
print(f"Excel文件数量：{len(files)}")
print(f"工作表总数量：{total_sheets}")
print(f"扫描报告：{report}")
